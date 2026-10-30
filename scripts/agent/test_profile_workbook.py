"""Generated test_only headerless workbooks; no source records or expected operational facts."""

import csv
import hashlib
from argparse import Namespace
from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook

from export_workbook_sheet import CsvRows, export
from profile_workbook import xlsx_profile
from read_workbook_windows import Refused, native

REPO = Path(__file__).resolve().parents[2]


def limits(path: Path, **extra) -> Namespace:
    return Namespace(services_geo_path=str(REPO / "services/geo"), file_path=str(path), rows=1000000,
                     record_bytes=1048576, member_bytes=7 * 1024 ** 3, ratio=100, bytes=1024 ** 3, **extra)


def profile(path: Path) -> list[dict]:
    workbook, _ = native(str(REPO / "services/geo"))
    with zipfile.ZipFile(path) as archive:
        return xlsx_profile(workbook, archive, limits(path))


@pytest.mark.parametrize("titles,width", [(0, 12), (1, 13), (2, 12), (3, 13)])
def test_headerless_keeps_every_data_row(tmp_path: Path, titles: int, width: int) -> None:
    path = tmp_path / "test_only.xlsx"
    book = Workbook()
    sheet = book.active
    for title in range(titles):
        sheet.append([f"test_only title {title}"])
    data = [["test_only text", "test_only label", *range(row, row + width - 2)] for row in range(125)]
    for row in data:
        sheet.append(row)
    book.save(path)
    original_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    result, = profile(path)
    assert result == {"name": "Sheet", "headerRow": None, "headerBasis": "absent", "titleRows": titles,
                      "firstDataRow": titles + 1, "occupiedColumns": list(range(1, width + 1)),
                      "rows": 125, "skip": None}
    output = tmp_path / "sheet.csv"
    settings = limits(path, sheet="Sheet", header=titles + 1, headerless=True,
                      columns=",".join(map(str, result["occupiedColumns"])))
    rows = CsvRows(output, settings)
    try:
        export(settings, rows)
        assert rows.records == 125
    finally:
        rows.target.close()
    with output.open(encoding="utf-8", newline="") as stream:
        written = list(csv.reader(stream))
    assert written == [[f"column_{index + 1}" for index in range(width)],
                       *[[str(value) for value in row] for row in data]]
    assert hashlib.sha256(output.read_bytes()).hexdigest() == rows.digest.hexdigest()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == original_hash


@pytest.mark.parametrize("shape", ["header", "ragged", "four_titles", "kind_change", "marked", "leading_blank"])
def test_uncertain_shapes_do_not_become_headerless(tmp_path: Path, shape: str) -> None:
    path = tmp_path / "test_only.xlsx"
    book = Workbook()
    sheet = book.active
    data = [["test_only text", "test_only label", *range(row, row + 10)] for row in range(5)]
    if shape == "header":
        sheet.append([f"stated_{column}" for column in range(12)])
    elif shape == "four_titles":
        for _ in range(4):
            sheet.append(["test_only title"])
    elif shape == "leading_blank":
        sheet.append([])
    if shape == "ragged":
        data[-1][3] = None
    elif shape == "kind_change":
        data[-1][3] = "test_only text"
    elif shape == "marked":
        data[-1][3] = "=1+1"
    for row in data:
        sheet.append(row)
    book.save(path)
    result, = profile(path)
    assert result["headerRow"] == (1 if shape == "header" else None)
    assert "headerBasis" not in result


def test_sparse_positions_are_compacted_without_exporting_titles(tmp_path: Path) -> None:
    path = tmp_path / "test_only.xlsx"
    book = Workbook()
    sheet = book.active
    for row in range(1, 4):
        sheet.cell(row, 2, "test_only text")
        sheet.cell(row, 4, row)
        sheet.cell(row, 7, row + 1)
        sheet.cell(row, 9, " ")
    book.save(path)
    result, = profile(path)
    assert result["occupiedColumns"] == [2, 4, 7]
    settings = limits(path, sheet="Sheet", header=1, headerless=True, columns="2,4,7")
    output = tmp_path / "sheet.csv"
    rows = CsvRows(output, settings)
    try:
        export(settings, rows)
    finally:
        rows.target.close()
    with output.open(encoding="utf-8", newline="") as stream:
        assert list(csv.reader(stream)) == [["column_1", "column_2", "column_3"],
                                            *[["test_only text", str(row), str(row + 1)] for row in range(1, 4)]]


def test_headerless_row_ceiling_includes_the_first_data_row(tmp_path: Path) -> None:
    settings = Namespace(header=1, headerless=True, columns="1,2", rows=2, record_bytes=1024, bytes=4096)
    rows = CsvRows(tmp_path / "sheet.csv", settings)
    try:
        rows.take(1, {1: ("text", None), 2: ("1", None)})
        rows.take(2, {1: ("text", None), 2: ("2", None)})
        assert rows.records == 2
        with pytest.raises(Refused) as caught:
            rows.take(3, {1: ("text", None), 2: ("3", None)})
        assert caught.value.code == "TABULAR_STREAM_ROWS"
    finally:
        rows.target.close()
